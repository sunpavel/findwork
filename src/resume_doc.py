#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Рендер структурированного резюме (из career_agent) в PDF и DOCX.

PDF  — fpdf2 + системный шрифт DejaVu (кириллица).
DOCX — python-docx (редактируемый Word-файл).

Структура резюме (dict):
  full_name, target_title, contacts{phone,email,telegram,location},
  profile, competencies[], experience[{company,title,period,context,
  responsibilities[],achievements[]}], education[], tools[], additional[]

Пакеты ставятся отдельно:  pip install fpdf2 python-docx
"""

from __future__ import annotations

import glob
import os
import re


def _find_fonts() -> tuple[str, str]:
    reg = next((p for p in (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans.ttf") if os.path.exists(p)), "")
    bold = next((p for p in (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf") if os.path.exists(p)), "")
    if not reg:
        hits = glob.glob("/usr/share/fonts/**/DejaVuSans.ttf", recursive=True)
        reg = hits[0] if hits else ""
    if not bold:
        hits = glob.glob("/usr/share/fonts/**/DejaVuSans-Bold.ttf", recursive=True)
        bold = hits[0] if hits else reg
    if not reg:
        raise RuntimeError("не найден шрифт DejaVuSans.ttf — установи: apt install fonts-dejavu-core")
    return reg, bold


def _contact_line(r: dict) -> str:
    c = r.get("contacts", {})
    bits = [c.get("location", ""), c.get("phone", ""), c.get("email", ""), c.get("telegram", "")]
    return "  |  ".join(b for b in bits if b)


def slug(s: str, default: str = "resume") -> str:
    s = re.sub(r"[^\w\s-]", "", (s or "")).strip()
    s = re.sub(r"[\s]+", "_", s)
    return s[:60] or default


# --- PDF ---------------------------------------------------------------------

def render_pdf(r: dict, path: str) -> str:
    from fpdf import FPDF  # noqa: PLC0415 — опциональная зависимость

    reg, bold = _find_fonts()
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_margins(15, 15, 15)
    pdf.add_page()
    pdf.add_font("dj", "", reg)
    pdf.add_font("dj", "B", bold)

    def header(title: str):
        pdf.ln(2)
        pdf.set_font("dj", "B", 11)
        pdf.set_text_color(20, 20, 20)
        pdf.cell(0, 7, title.upper(), new_x="LMARGIN", new_y="NEXT")
        y = pdf.get_y()
        pdf.set_draw_color(180, 180, 180)
        pdf.line(15, y, 195, y)
        pdf.ln(1)
        pdf.set_text_color(0, 0, 0)

    def cell(text: str, size=10, style=""):
        pdf.set_font("dj", style, size)
        pdf.multi_cell(0, 5, text, new_x="LMARGIN", new_y="NEXT")

    body = cell

    def bullets(items):
        pdf.set_font("dj", "", 10)
        for it in items:
            if it:
                pdf.multi_cell(0, 5, f"- {it}", new_x="LMARGIN", new_y="NEXT")

    # Шапка
    pdf.set_font("dj", "B", 17)
    pdf.cell(0, 9, r.get("full_name", ""), new_x="LMARGIN", new_y="NEXT")
    if r.get("target_title"):
        pdf.set_font("dj", "B", 12)
        pdf.set_text_color(70, 70, 70)
        pdf.cell(0, 7, r["target_title"], new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(0, 0, 0)
    line = _contact_line(r)
    if line:
        pdf.set_font("dj", "", 9)
        pdf.multi_cell(0, 5, line, new_x="LMARGIN", new_y="NEXT")

    if r.get("profile"):
        header("Профессиональный профиль")
        body(r["profile"])
    if r.get("competencies"):
        header("Ключевые компетенции")
        body(" · ".join(r["competencies"]))
    if r.get("experience"):
        header("Опыт работы")
        for e in r["experience"]:
            pdf.ln(1)
            pdf.set_font("dj", "B", 10.5)
            pdf.multi_cell(0, 5, e.get("company", ""), new_x="LMARGIN", new_y="NEXT")
            sub = " · ".join(x for x in (e.get("title", ""), e.get("period", ""),
                                         e.get("context", "")) if x)
            if sub:
                pdf.set_font("dj", "", 9)
                pdf.set_text_color(90, 90, 90)
                pdf.multi_cell(0, 5, sub, new_x="LMARGIN", new_y="NEXT")
                pdf.set_text_color(0, 0, 0)
            bullets(e.get("responsibilities", []))
            bullets(e.get("achievements", []))
    if r.get("education"):
        header("Образование")
        bullets(r["education"])
    if r.get("tools"):
        header("Инструменты")
        body(" · ".join(r["tools"]))
    if r.get("additional"):
        header("Дополнительно")
        bullets(r["additional"])

    pdf.output(path)
    return path


# --- DOCX --------------------------------------------------------------------

def render_docx(r: dict, path: str) -> str:
    from docx import Document  # noqa: PLC0415
    from docx.shared import Pt, RGBColor  # noqa: PLC0415

    doc = Document()
    normal = doc.styles["Normal"].font
    normal.name = "Calibri"
    normal.size = Pt(10)

    def header(title: str):
        p = doc.add_paragraph()
        run = p.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)

    name = doc.add_paragraph()
    run = name.add_run(r.get("full_name", ""))
    run.bold = True
    run.font.size = Pt(17)
    if r.get("target_title"):
        t = doc.add_paragraph()
        tr = t.add_run(r["target_title"])
        tr.bold = True
        tr.font.size = Pt(12)
        tr.font.color.rgb = RGBColor(0x46, 0x46, 0x46)
    line = _contact_line(r)
    if line:
        cp = doc.add_paragraph()
        cp.add_run(line).font.size = Pt(9)

    if r.get("profile"):
        header("Профессиональный профиль")
        doc.add_paragraph(r["profile"])
    if r.get("competencies"):
        header("Ключевые компетенции")
        doc.add_paragraph(" · ".join(r["competencies"]))
    if r.get("experience"):
        header("Опыт работы")
        for e in r["experience"]:
            p = doc.add_paragraph()
            p.add_run(e.get("company", "")).bold = True
            sub = " · ".join(x for x in (e.get("title", ""), e.get("period", ""),
                                         e.get("context", "")) if x)
            if sub:
                sp = doc.add_paragraph()
                srun = sp.add_run(sub)
                srun.italic = True
                srun.font.size = Pt(9)
                srun.font.color.rgb = RGBColor(0x5A, 0x5A, 0x5A)
            for it in e.get("responsibilities", []) + e.get("achievements", []):
                if it:
                    doc.add_paragraph(it, style="List Bullet")
    if r.get("education"):
        header("Образование")
        for it in r["education"]:
            doc.add_paragraph(it, style="List Bullet")
    if r.get("tools"):
        header("Инструменты")
        doc.add_paragraph(" · ".join(r["tools"]))
    if r.get("additional"):
        header("Дополнительно")
        for it in r["additional"]:
            doc.add_paragraph(it, style="List Bullet")

    doc.save(path)
    return path


def render_both(r: dict, out_dir: str) -> tuple[str, str]:
    """Рендерит PDF и DOCX в out_dir. Возвращает (pdf_path, docx_path)."""
    base = "Резюме_" + slug(r.get("target_title") or r.get("full_name"))
    pdf = render_pdf(r, os.path.join(out_dir, base + ".pdf"))
    docx = render_docx(r, os.path.join(out_dir, base + ".docx"))
    return pdf, docx
